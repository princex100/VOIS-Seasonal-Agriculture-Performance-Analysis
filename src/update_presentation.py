"""
Automated PowerPoint Presentation Deck Generator
VOIS AICTE Internship (Batch 1 2026-2027) Major Project
Title: Seasonal Agriculture Performance Analysis

Module: src/update_presentation.py
Generates: docs/VOIS_Major_Project_Final_Submission.pptx (14 Slides)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# Color Constants & Palette
# -------------------------------------------------------------
COLOR_RED = RGBColor(230, 0, 0)        # #E60000 (Vodafone/VOIS Red)
COLOR_DARK = RGBColor(28, 28, 28)      # #1C1C1C
COLOR_SLATE = RGBColor(60, 64, 67)     # #3C4043
COLOR_LIGHT_BG = RGBColor(246, 248, 250) # #F6F8FA
COLOR_CARD_BG = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(220, 224, 230)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_TEAL = RGBColor(0, 139, 139)
COLOR_GOLD = RGBColor(230, 149, 0)

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="VOIS AICTE INTERNSHIP - MAJOR PROJECT"):
    # Header Banner Shape
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_DARK
    top_bar.line.fill.background()

    # Red accent line below header
    red_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.06))
    red_line.fill.solid()
    red_line.fill.fore_color.rgb = COLOR_RED
    red_line.line.fill.background()

    # Category Text
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_RED
    p_cat.font.name = 'Arial'

    # Title Text
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.font.name = 'Arial'

def add_footer(slide, slide_num, total_slides=14):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"VOIS AICTE Internship (Batch 1 2026-2027) | Seasonal Agriculture Performance Analysis | Slide {slide_num} of {total_slides}"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(128, 134, 139)
    p.font.name = 'Arial'

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_DARK)

    # Decorative red backdrop bar
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.18), Inches(5.0))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_RED
    accent_bar.line.fill.background()

    # Main Title Box
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(2.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "VOIS AICTE INTERNSHIP (BATCH 1: 2026-2027) | MAJOR PROJECT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED
    p0.font.name = 'Arial'

    p1 = tf1.add_paragraph()
    p1.text = "Seasonal Agriculture Performance Analysis"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.font.name = 'Arial'
    p1.space_before = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Empirical Data Analytics, Environmental Risk Modeling & Agritech Strategic Optimization"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(200, 205, 210)
    p2.font.name = 'Arial'
    p2.space_before = Pt(8)

    # Student Details Card
    card_info = add_card(slide1, Inches(1.2), Inches(4.0), Inches(10.8), Inches(2.2), bg_color=RGBColor(38, 42, 46), border_color=COLOR_RED)
    
    tf_info = card_info.text_frame
    tf_info.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_info.word_wrap = True

    p_info_title = tf_info.paragraphs[0]
    p_info_title.text = "PROJECT SUBMISSION CREDENTIALS"
    p_info_title.font.size = Pt(11)
    p_info_title.font.bold = True
    p_info_title.font.color.rgb = COLOR_RED
    p_info_title.font.name = 'Arial'

    details_text = [
        ("Student Name:", "[Student Name]"),
        ("College Name:", "[College Name]"),
        ("AICTE Student ID:", "[AICTE STU ID]"),
        ("Internship Track:", "Data Visualization & Analytics | Vodafone Intelligent Solutions (_VOIS)"),
        ("Academic Batch:", "Batch 1 (2026-2027)")
    ]

    for label, val in details_text:
        p_row = tf_info.add_paragraph()
        p_row.text = f"•  {label} {val}"
        p_row.font.size = Pt(12)
        p_row.font.color.rgb = COLOR_WHITE
        p_row.font.name = 'Arial'
        p_row.space_before = Pt(3)

    add_footer(slide1, 1)

    # =========================================================================
    # SLIDE 2: Problem Statement
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_LIGHT_BG)
    add_header(slide2, "Problem Statement: Challenges in Indian Seasonal Agriculture")

    # Card 1: Seasonal Vulnerabilities
    card1 = add_card(slide2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    p_c1 = tf_c1.paragraphs[0]
    p_c1.text = "🚨 Core Problem & Seasonal Vulnerabilities"
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_RED

    points_c1 = [
        "Unpredictable Seasonal Transitions: Extreme meteorological variations between Monsoon (Kharif), Winter (Rabi), and Summer (Zaid) create volatile yields.",
        "Monsoon Over-Reliance: Over 60% of smallholders lack dynamic adaptive planning, exposing crops to sudden moisture deficits or severe humidity-induced pest surges.",
        "Economic Margin Erosion: High input costs (power, water pumping, fertilizers) in summer cycles (Zaid) frequently exceed harvest market revenues."
    ]
    for pt in points_c1:
        p = tf_c1.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(12)

    # Card 2: Knowledge Gaps & Solution
    card2 = add_card(slide2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    p_c2 = tf_c2.paragraphs[0]
    p_c2.text = "💡 Research Scope & Empirical Objectives"
    p_c2.font.size = Pt(16)
    p_c2.font.bold = True
    p_c2.font.color.rgb = COLOR_DARK

    points_c2 = [
        "Eliminate Empirical Intuition: Replace traditional unquantified farming decisions with rigorous statistical and financial modeling.",
        "Resource Mismatch Rectification: Quantify irrigation inefficiencies (Flood vs. Drip) to prevent groundwater squandering and input over-application.",
        "Evidence-Based Agritech Planning: Deliver clear crop-season affinity matrices and actionable decision thresholds for policy makers and farming communities."
    ]
    for pt in points_c2:
        p = tf_c2.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_SLATE
        p.space_before = Pt(12)

    add_footer(slide2, 2)

    # =========================================================================
    # SLIDE 3: Project Description
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_LIGHT_BG)
    add_header(slide3, "Project Description & Dataset Architecture")

    # Metrics Overview Cards
    stat_boxes = [
        ("4,000", "Farm Samples Analyzed", COLOR_RED),
        ("8 States", "Pan-India Agricultural Coverage", COLOR_DARK),
        ("3 Seasons", "Kharif, Rabi, and Zaid Cycles", COLOR_TEAL),
        ("28 Features", "Environmental, Soil & Economics", COLOR_GOLD)
    ]
    for i, (val, lbl, col) in enumerate(stat_boxes):
        x_pos = Inches(0.8 + i * 2.98)
        c = add_card(slide3, x_pos, Inches(1.4), Inches(2.8), Inches(1.3))
        tf = c.text_frame
        tf.word_wrap = True
        p_v = tf.paragraphs[0]
        p_v.text = val
        p_v.font.size = Pt(22)
        p_v.font.bold = True
        p_v.font.color.rgb = col
        p_v.alignment = PP_ALIGN.CENTER
        
        p_l = tf.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(9.5)
        p_l.font.color.rgb = COLOR_SLATE
        p_l.alignment = PP_ALIGN.CENTER

    # Detailed Scope Card
    c_desc = add_card(slide3, Inches(0.8), Inches(2.9), Inches(11.733), Inches(3.8))
    tf_d = c_desc.text_frame
    tf_d.word_wrap = True
    p_d0 = tf_d.paragraphs[0]
    p_d0.text = "Comprehensive Agronomic & Economic Analytical Dimensions"
    p_d0.font.size = Pt(15)
    p_d0.font.bold = True
    p_d0.font.color.rgb = COLOR_DARK

    desc_bullets = [
        "Geographic & Soil Diversity: Granular observations across Andhra Pradesh, Gujarat, Karnataka, Madhya Pradesh, Maharashtra, Punjab, Tamil Nadu, and Telangana covering pH (5.2 - 8.4), Soil Moisture, and N-P-K nutrient profiles.",
        "Meteorological & Irrigation Inputs: Systematic telemetry including Seasonal Rainfall (mm), Temperature (°C), Humidity (%), Sunlight Hours, and Water Volumes across 4 irrigation modes (Drip, Sprinkler, Rainfed, Flood).",
        "Financial & Productivity Metrics: Full micro-economic auditing computing Production (Tonnes), Cost (₹), Revenue (₹), Net Profit (₹), Profit/Ha, and Water Productivity ($t/1,000 m^3$).",
        "Risk Profiling: Quantitative assessment of chemical load (Pesticides, Fertilizers, Seed Quality) against Disease and Pest Vulnerability risk percentages."
    ]
    for b in desc_bullets:
        p = tf_d.add_paragraph()
        p.text = f"•  {b}"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_SLATE
        p.space_before = Pt(8)

    add_footer(slide3, 3)

    # =========================================================================
    # SLIDE 4: End Users & Stakeholders
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_LIGHT_BG)
    add_header(slide4, "Target End Users & Stakeholder Value Proposition")

    users = [
        ("👨‍🌾 Smallholder Farmers",
         "Operational Decisions",
         ["Select optimal crop-season combinations to maximize net margins.",
          "Adopt Drip/Sprinkler methods to cut water waste and double ROI.",
          "Prevent pest outbreaks by timing chemical sprays with humidity alerts."],
         COLOR_RED),
        ("🏦 Agri-Fintech & Insurers",
         "Risk & Credit Underwriting",
         ["Structure parametric insurance based on seasonal rainfall thresholds.",
          "Assess farmer creditworthiness using empirical profit-per-hectare data.",
          "Reduce loan default rates through automated seasonal risk scoring."],
         COLOR_DARK),
        ("🚚 Supply Chain & FMCG",
         "Procurement & Inventory",
         ["Forecast harvest volume across Kharif and Rabi with high fidelity.",
          "Optimize warehousing capacity before seasonal harvest arrivals.",
          "Minimize supply shocks for agro-processing commodities (Sugarcane, Chilli)."],
         COLOR_TEAL),
        ("🏛️ Extension Officers & Policy",
         "Agritech Policy Implementation",
         ["Direct government subsidies towards high-efficiency micro-irrigation.",
          "Target IPM (Integrated Pest Management) campaigns in high-risk Kharif zones.",
          "Deploy regional soil health interventions based on N-P-K deficiency maps."],
         COLOR_GOLD)
    ]

    for i, (title, subtitle, bullets, accent_col) in enumerate(users):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(1.5 + (i // 2) * 2.7)
        c = add_card(slide4, x, y, Inches(5.75), Inches(2.55))
        tf = c.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = accent_col

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle.upper()
        p_sub.font.size = Pt(9.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_SLATE

        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_DARK
            p.space_before = Pt(3)

    add_footer(slide4, 4)

    # =========================================================================
    # SLIDE 5: Technology Used
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_LIGHT_BG)
    add_header(slide5, "Technology Stack & Analytical Pipeline")

    tech_categories = [
        ("Python 3.10+ & Data Engine", ["Pandas: High-performance data manipulation & aggregation", "NumPy: Numerical vectorization & matrix operations"], COLOR_RED),
        ("Statistical & Inferential Modeling", ["Scipy.stats: ANOVA ($F$-statistic) and Kruskal-Wallis ($H$-test)", "Statsmodels: Linear OLS regression and trend analysis"], COLOR_DARK),
        ("Publication Visualization", ["Matplotlib: High-resolution dual-axis and grouped layouts", "Seaborn: Heatmap matrices and statistical regression plotting"], COLOR_TEAL),
        ("Automation & Deliverables", ["python-pptx: Programmatic corporate deck generation", "nbformat & Jupyter: Documented reproducible research notebook", "Git & GitHub: Version control & public repository index"], COLOR_GOLD)
    ]

    for i, (cat, items, col) in enumerate(tech_categories):
        y = Inches(1.5 + i * 1.35)
        c = add_card(slide5, Inches(0.8), y, Inches(11.733), Inches(1.2))
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"⚙️ {cat}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"   •  {item}"
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = COLOR_DARK
            p_item.space_before = Pt(2)

    add_footer(slide5, 5)

    # =========================================================================
    # HELPER FOR RESULT SLIDES (SLIDES 6 - 10)
    # =========================================================================
    def build_result_slide(slide_num, title, img_path, takeaways):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, COLOR_LIGHT_BG)
        add_header(slide, title, category_text=f"RESULT & EMPIRICAL FINDINGS (CHART {slide_num-5} OF 5)")

        # Left Image Placement (High DPI PNG)
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.45), Inches(6.8), Inches(5.2))
        else:
            c_placeholder = add_card(slide, Inches(0.8), Inches(1.45), Inches(6.8), Inches(5.2))
            tf = c_placeholder.text_frame
            p = tf.paragraphs[0]
            p.text = f"[Image not found: {img_path}]"

        # Right Key Takeaways Card
        c_takeaways = add_card(slide, Inches(7.8), Inches(1.45), Inches(4.733), Inches(5.2))
        tf_t = c_takeaways.text_frame
        tf_t.word_wrap = True

        p_head = tf_t.paragraphs[0]
        p_head.text = "📊 Key Analytical Insights"
        p_head.font.size = Pt(15)
        p_head.font.bold = True
        p_head.font.color.rgb = COLOR_RED

        for idx, (head, body) in enumerate(takeaways, 1):
            p_h = tf_t.add_paragraph()
            p_h.text = f"{idx}. {head}"
            p_h.font.size = Pt(12)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_DARK
            p_h.space_before = Pt(12)

            p_b = tf_t.add_paragraph()
            p_b.text = body
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = COLOR_SLATE
            p_b.space_before = Pt(2)

        add_footer(slide, slide_num)

    # =========================================================================
    # SLIDE 6: Result 1 - Seasonal Yield vs. Rainfall
    # =========================================================================
    build_result_slide(
        slide_num=6,
        title="Result 1: Seasonal Agronomic Dynamics (Yield vs. Rainfall)",
        img_path="results/01_seasonal_yield_rainfall.png",
        takeaways=[
            ("Monsoon-Driven Crop Yields",
             "Kharif yields peak at 5.63 Tonnes/Ha driven by heavy monsoon rainfall (mean 852.1 mm), compared to 5.09 t/Ha in Rabi and 4.64 t/Ha in Zaid."),
            ("Precipitation Gradient Across Cycles",
             "Rainfall drops steeply by 48.8% in Rabi (435.9 mm) and 64.9% in Zaid (299.2 mm), necessitating supplementary artificial irrigation."),
            ("Statistically Significant Variation",
             "Kruskal-Wallis testing confirms highly significant seasonal differences in crop yield distributions (H = 70.26, p = 5.52 × 10⁻¹⁶).")
        ]
    )

    # =========================================================================
    # SLIDE 7: Result 2 - Irrigation Efficiency vs Profit
    # =========================================================================
    build_result_slide(
        slide_num=7,
        title="Result 2: Irrigation Method Efficiency vs. Net Profitability",
        img_path="results/02_irrigation_profit_efficiency.png",
        takeaways=[
            ("Drip Irrigation Financial Premium",
             "Drip irrigation achieves the highest average net profit of ₹25,544/Ha (₹218,659/farm), nearly triple that of traditional flood methods."),
            ("High Water Productivity",
             "Drip delivers 6.26 t/1,000 m³ water efficiency, ensuring maximum biological crop output per unit volume of water consumed."),
            ("Flood Irrigation Inefficiency",
             "Flood irrigation proves lowest in profitability (₹9,005/Ha) and water efficiency (3.44 t/1,000 m³), squandering massive ground table reserves.")
        ]
    )

    # =========================================================================
    # SLIDE 8: Result 3 - Economic Returns across Seasons
    # =========================================================================
    build_result_slide(
        slide_num=8,
        title="Result 3: Seasonal Macroeconomic Breakdown (Cost, Revenue, Profit)",
        img_path="results/03_economic_returns_season.png",
        takeaways=[
            ("Kharif High-Profit Regime",
             "Kharif generates ₹711k in revenue against ₹532k in cost, producing the highest seasonal profit margin of ₹179k per farm (₹21,973/Ha)."),
            ("Rabi Economic Stability",
             "Rabi produces ₹602k revenue and ₹514k cost, delivering stable positive net earnings of ₹88k per farm (₹10,429/Ha)."),
            ("Zaid Financial Inversion Deficit",
             "Zaid farms suffer an average net deficit of -₹24k (-₹2,519/Ha) as high summer pumping/labor costs (₹544k) outpace realized revenues (₹520k). ANOVA F = 34.40, p = 1.54 × 10⁻¹⁵.")
        ]
    )

    # =========================================================================
    # SLIDE 9: Result 4 - Agrochemicals vs. Pest Risk
    # =========================================================================
    build_result_slide(
        slide_num=9,
        title="Result 4: Agrochemical Application vs. Pest/Disease Vulnerability",
        img_path="results/04_pesticide_fertilizer_pest_risk.png",
        takeaways=[
            ("Environmental Climatological Driver",
             "Pest vulnerability surges to 54.47% in Kharif vs. 40.48% in Rabi and 38.22% in Zaid, driven by monsoon atmospheric humidity > 75%."),
            ("Weak Chemical Dosage Correlation",
             "Pesticide dosage (r = 0.015, p = 0.35) and fertilizer dosage (r = 0.033, p = 0.038) demonstrate negligible linear correlation with pest risk reduction."),
            ("Integrated Pest Management (IPM) Need",
             "Uncalibrated chemical over-application fails to mitigate risk; preventive telemetry and biological pest controls are essential.")
        ]
    )

    # =========================================================================
    # SLIDE 10: Result 5 - Crop Performance Matrix
    # =========================================================================
    build_result_slide(
        slide_num=10,
        title="Result 5: Seasonal Crop Performance Matrix (Profit & Productivity)",
        img_path="results/05_crop_performance_matrix.png",
        takeaways=[
            ("High-Margin Cash Crops",
             "Sugarcane and Chilli generate superior profitability (₹60,000 - ₹150,000/Ha) and tonnage (Sugarcane: 40 - 75 t/Ha) across Kharif and Rabi."),
            ("Food Security Staples",
             "Rice (2.5 - 4.5 t/Ha) and Wheat (2.0 - 4.0 t/Ha) deliver steady staple productivity with low volatility across Kharif and Rabi."),
            ("Zaid Crop Selection Guidelines",
             "Summer farming should prioritize drought-tolerant pulses and short-duration oilseeds (Groundnut) supported by drip systems.")
        ]
    )

    # =========================================================================
    # SLIDE 11: Future Scope & Roadmap
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, COLOR_LIGHT_BG)
    add_header(slide11, "Future Scope & Agritech Innovation Roadmap")

    roadmap = [
        ("🤖 Machine Learning Yield Prediction",
         "Deploy XGBoost, LightGBM, and Random Forest ensemble regressors to predict micro-climate farm yields with R² > 0.90.",
         COLOR_RED),
        ("📡 Real-Time IoT Soil Sensor Integration",
         "Ingest telemetry from ground-level N-P-K, pH, and capacitive soil moisture probes for automated variable-rate drip fertigation.",
         COLOR_DARK),
        ("🌦️ Micro-Climate Crop Advisory Systems",
         "AI-driven hyper-local mobile advisory delivering real-time pest alerts and market price intelligence directly to smallholders.",
         COLOR_TEAL),
        ("🛰️ Satellite Remote Sensing (NDVI/EVI)",
         "Incorporate Sentinel-2 and Landsat optical/SAR imagery to monitor crop canopy health and stress index in real-time.",
         COLOR_GOLD)
    ]

    for i, (title, desc, col) in enumerate(roadmap):
        y = Inches(1.5 + i * 1.35)
        c = add_card(slide11, Inches(0.8), y, Inches(11.733), Inches(1.2))
        tf = c.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_DARK
        p_d.space_before = Pt(3)

    add_footer(slide11, 11)

    # =========================================================================
    # SLIDE 12: GitHub Repository & Google Colab Links
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, COLOR_LIGHT_BG)
    add_header(slide12, "Project Repository & Google Colab Interactive Links")

    # Card 1: GitHub Repository Link Card (Left)
    c_gh = add_card(slide12, Inches(0.8), Inches(1.4), Inches(5.75), Inches(2.2), bg_color=COLOR_DARK, border_color=COLOR_RED)
    tf_gh = c_gh.text_frame
    tf_gh.word_wrap = True
    p_gh0 = tf_gh.paragraphs[0]
    p_gh0.text = "🐙 OFFICIAL GITHUB REPOSITORY"
    p_gh0.font.size = Pt(11)
    p_gh0.font.bold = True
    p_gh0.font.color.rgb = COLOR_RED

    p_gh1 = tf_gh.add_paragraph()
    p_gh1.text = "https://github.com/princex100/VOIS-Seasonal-Agriculture-Performance-Analysis"
    p_gh1.font.size = Pt(12)
    p_gh1.font.bold = True
    p_gh1.font.color.rgb = COLOR_WHITE
    p_gh1.space_before = Pt(4)

    p_gh2 = tf_gh.add_paragraph()
    p_gh2.text = "• Complete production codebase, datasets, and presentation deck.\n• Full Git commit history & structured directory hierarchy."
    p_gh2.font.size = Pt(10)
    p_gh2.font.color.rgb = RGBColor(200, 205, 210)
    p_gh2.space_before = Pt(4)

    # Card 2: Google Colab 1-Click Notebook Link Card (Right)
    c_colab = add_card(slide12, Inches(6.78), Inches(1.4), Inches(5.75), Inches(2.2), bg_color=COLOR_DARK, border_color=COLOR_GOLD)
    tf_colab = c_colab.text_frame
    tf_colab.word_wrap = True
    p_cl0 = tf_colab.paragraphs[0]
    p_cl0.text = "⚡ GOOGLE COLAB INTERACTIVE NOTEBOOK"
    p_cl0.font.size = Pt(11)
    p_cl0.font.bold = True
    p_cl0.font.color.rgb = COLOR_GOLD

    p_cl1 = tf_colab.add_paragraph()
    p_cl1.text = "https://colab.research.google.com/github/princex100/VOIS-Seasonal-Agriculture-Performance-Analysis/blob/main/notebooks/Seasonal_Agriculture_Performance_Analysis.ipynb"
    p_cl1.font.size = Pt(10.5)
    p_cl1.font.bold = True
    p_cl1.font.color.rgb = COLOR_WHITE
    p_cl1.space_before = Pt(4)

    p_cl2 = tf_colab.add_paragraph()
    p_cl2.text = "• 1-Click Cloud Execution: Run end-to-end EDA and statistical ANOVA directly in your browser.\n• Zero setup required: Auto-fetches dataset and renders interactive charts."
    p_cl2.font.size = Pt(10)
    p_cl2.font.color.rgb = RGBColor(200, 205, 210)
    p_cl2.space_before = Pt(4)

    # Repository Structure Card (Bottom)
    c_tree = add_card(slide12, Inches(0.8), Inches(3.8), Inches(11.733), Inches(3.0))
    tf_t = c_tree.text_frame
    tf_t.word_wrap = True
    p_t0 = tf_t.paragraphs[0]
    p_t0.text = "📁 Structured Deliverables & Reproduction Directory Index"
    p_t0.font.size = Pt(13)
    p_t0.font.bold = True
    p_t0.font.color.rgb = COLOR_DARK

    tree_items = [
        ("data/", "Contains raw CSV (4,000 records) and cleaned dataset with stratified median imputation."),
        ("notebooks/", "Seasonal_Agriculture_Performance_Analysis.ipynb - Google Colab compatible research notebook."),
        ("src/", "analyze_and_visualize.py (Data & EDA pipeline), update_presentation.py (Deck builder)."),
        ("results/", "5 High-resolution publication charts (.png) and statistical_summary.json metrics."),
        ("docs/", "VOIS_Major_Project_Final_Submission.pptx - 14-slide executive submission presentation.")
    ]
    for folder, desc in tree_items:
        p = tf_t.add_paragraph()
        p.text = f"•  {folder:<15} : {desc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_SLATE
        p.space_before = Pt(2)

    add_footer(slide12, 12)

    # =========================================================================
    # SLIDE 13: Certificate Slide
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, COLOR_LIGHT_BG)
    add_header(slide13, "VOIS Data Visualization Course Completion Certificate")

    # Outer Certificate Placement Frame
    c_frame = add_card(slide13, Inches(1.8), Inches(1.5), Inches(9.733), Inches(5.1), bg_color=COLOR_WHITE, border_color=COLOR_RED)
    tf_f = c_frame.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_f.word_wrap = True

    p_banner = tf_f.paragraphs[0]
    p_banner.text = "[Insert VOIS Data Visualization Certificate Here]"
    p_banner.font.size = Pt(22)
    p_banner.font.bold = True
    p_banner.font.color.rgb = COLOR_RED
    p_banner.alignment = PP_ALIGN.CENTER

    p_sub = tf_f.add_paragraph()
    p_sub.text = "Paste your official Vodafone Intelligent Solutions (_VOIS) Data Visualization course certificate image here prior to final submission."
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = COLOR_SLATE
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(10)

    add_footer(slide13, 13)

    # =========================================================================
    # SLIDE 14: Thank You & Contact
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, COLOR_DARK)

    # Center box
    c_ty = slide14.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.8))
    tf_ty = c_ty.text_frame
    tf_ty.word_wrap = True

    p0 = tf_ty.paragraphs[0]
    p0.text = "Thank You!"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_WHITE
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf_ty.add_paragraph()
    p1.text = "VOIS AICTE INTERNSHIP (BATCH 1 2026-2027) - MAJOR PROJECT"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_RED
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(10)

    p2 = tf_ty.add_paragraph()
    p2.text = "Seasonal Agriculture Performance Analysis"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(220, 225, 230)
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)

    p3 = tf_ty.add_paragraph()
    p3.text = "Contact & Student Information:\nStudent: [Student Name]  |  Email: [Your Email]  |  LinkedIn: [Your LinkedIn Profile]\nAICTE Student ID: [AICTE STU ID]  |  College: [College Name]"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(180, 185, 190)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(25)

    add_footer(slide14, 14)

    # Save output
    output_path = "docs/VOIS_Major_Project_Final_Submission.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path} (14 Slides, 16:9 Widescreen)")

if __name__ == "__main__":
    create_presentation()
