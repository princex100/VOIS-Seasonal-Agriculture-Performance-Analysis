import os
from pptx import Presentation

prs = Presentation('docs/VOIS_Major_Project_Final_Submission.pptx')
print(f"Total slides: {len(prs.slides)}")
print(f"Dimensions: {prs.slide_width.inches} x {prs.slide_height.inches}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s in slide.shapes:
        s_type = 'Placeholder' if s.is_placeholder else ('TextFrame' if s.has_text_frame else str(s.shape_type))
        ph_type = s.placeholder_format.type if s.is_placeholder else 'None'
        text_content = []
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                text_content.append(p.text)
        text_str = " | ".join(text_content)
        print(f"Shape: [{s.name}] | Type: {s_type} (PH Type: {ph_type}) | Pos: (L={s.left.inches:.2f}, T={s.top.inches:.2f}, W={s.width.inches:.2f}, H={s.height.inches:.2f})")
        if text_str:
            print(f"   Text: {text_str}")
