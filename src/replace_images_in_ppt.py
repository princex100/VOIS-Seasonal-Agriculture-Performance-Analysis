"""
Replace chart images in the presentation with new images from Google Colab (vois1.png to vois5.png).
Preserves all user-entered details (name, college, AICTE ID, custom text).
"""

import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt

def replace_images():
    # 1. Update results directory images
    img_map = {
        'vois1.png': ('results/01_seasonal_yield_rainfall.png', 5), # Slide 6
        'vois2.png': ('results/02_irrigation_profit_efficiency.png', 6), # Slide 7
        'vois3.png': ('results/03_economic_returns_season.png', 7), # Slide 8
        'vois4.png': ('results/04_pesticide_fertilizer_pest_risk.png', 8), # Slide 9
        'vois5.png': ('results/05_crop_performance_matrix.png', 9), # Slide 10
    }

    for src_img, (dest_img, _) in img_map.items():
        if os.path.exists(src_img):
            shutil.copy2(src_img, dest_img)
            print(f"Copied {src_img} -> {dest_img}")
        else:
            print(f"Warning: {src_img} not found!")

    # 2. Open user's presentation
    ppt_source = "docs/vois-finalppt.pptx"
    if not os.path.exists(ppt_source):
        ppt_source = "docs/VOIS_Major_Project_Final_Submission.pptx"

    prs = Presentation(ppt_source)
    print(f"Loaded presentation: {ppt_source} ({len(prs.slides)} slides)")

    for src_img, (_, slide_idx) in img_map.items():
        if not os.path.exists(src_img):
            continue

        slide = prs.slides[slide_idx]
        
        # Identify and remove previous chart image shape
        shapes_to_delete = []
        for s in slide.shapes:
            if s.shape_type == 13: # Picture
                # Only remove main chart picture (not small bottom footer logo)
                if s.width.inches > 3.0 and s.top.inches < 6.0:
                    shapes_to_delete.append(s)

        for s in shapes_to_delete:
            sp_elem = s._element
            sp_elem.getparent().remove(sp_elem)
            print(f"Slide {slide_idx+1}: Removed old picture {s.name}")

        # Add the new picture from Colab
        new_pic = slide.shapes.add_picture(
            src_img,
            Inches(0.8),
            Inches(1.5),
            Inches(6.8),
            Inches(5.1)
        )
        print(f"Slide {slide_idx+1}: Inserted new image {src_img} -> {new_pic.name}")

    # 3. Save to presentation locations
    save_targets = [
        "docs/vois-finalppt.pptx",
        "docs/VOIS_Major_Project_Final_Submission.pptx",
        "C:/Users/LENOVO/OneDrive/Desktop/vois-finalppt.pptx"
    ]

    for target in save_targets:
        try:
            prs.save(target)
            print(f"Successfully saved updated presentation to: {target}")
        except Exception as e:
            print(f"Note: Could not save to {target}: {e}")

if __name__ == "__main__":
    replace_images()
